"""Экспорт отчётов."""
import csv
import io
from datetime import datetime
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from ..database import get_db
from ..models import (
    User, Project, AccessPoint, ClientDevice, PlanElement,
    Measurement, AuditLog, SSIDProfile,
)
from ..deps import require_role
from ..physics import coverage_statistics, analyze_channels

router = APIRouter()


def _load_aps(db, project_id):
    aps = db.query(AccessPoint).filter(AccessPoint.project_id == project_id).all()
    return [{"id": a.id, "name": a.name, "model": a.model,
             "x": a.x, "y": a.y, "tx_power_dbm": a.tx_power_dbm,
             "antenna_gain": a.antenna_gain, "band": a.band, "channel": a.channel}
            for a in aps]


def _load_walls(db, project_id):
    els = db.query(PlanElement).filter(
        PlanElement.project_id == project_id,
        PlanElement.type.in_(["wall", "door", "window"])
    ).all()
    return [{"x1": el.x1, "y1": el.y1, "x2": el.x2, "y2": el.y2,
             "material": el.material or "concrete", "type": el.type}
            for el in els]


@router.get("/{project_id}/csv")
def export_csv(project_id: int, db: Session = Depends(get_db),
               user: User = Depends(require_role("admin", "analyst"))):
    p = db.query(Project).get(project_id)
    if not p:
        raise HTTPException(status_code=404, detail="Проект не найден")
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(["# Отчёт о проекте", p.name])
    w.writerow(["Размер", f"{p.width_m}×{p.height_m} м"])
    w.writerow(["Дата", datetime.utcnow().isoformat()])
    w.writerow([])
    w.writerow(["# Точки доступа"])
    w.writerow(["id", "name", "model", "x_m", "y_m", "tx_dbm",
                "gain_dbi", "band", "channel", "ssid"])
    for ap in db.query(AccessPoint).filter(AccessPoint.project_id == project_id).all():
        w.writerow([ap.id, ap.name, ap.model or "", ap.x, ap.y,
                    ap.tx_power_dbm, ap.antenna_gain, ap.band,
                    ap.channel, ap.ssid_type])
    w.writerow([])
    w.writerow(["# Устройства"])
    w.writerow(["id", "name", "type", "x_m", "y_m", "required_rssi", "required_speed"])
    for d in db.query(ClientDevice).filter(ClientDevice.project_id == project_id).all():
        w.writerow([d.id, d.name, d.type, d.x, d.y,
                    d.required_rssi, d.required_speed])
    w.writerow([])
    w.writerow(["# Стены"])
    w.writerow(["id", "type", "material", "x1", "y1", "x2", "y2"])
    for el in db.query(PlanElement).filter(
        PlanElement.project_id == project_id,
        PlanElement.type.in_(["wall", "door", "window"])
    ).all():
        w.writerow([el.id, el.type, el.material or "",
                    el.x1, el.y1, el.x2, el.y2])
    w.writerow([])
    w.writerow(["# Измерения"])
    w.writerow(["id", "name", "x_m", "y_m", "rssi", "snr",
                "interference", "mode", "measured_at"])
    for m in db.query(Measurement).filter(Measurement.project_id == project_id).all():
        w.writerow([m.id, m.name or "", m.x, m.y, m.rssi,
                    m.snr or "", m.interference or "", m.mode,
                    m.measured_at.isoformat() if m.measured_at else ""])
    w.writerow([])
    w.writerow(["# Статистика покрытия"])
    aps = _load_aps(db, project_id)
    walls = _load_walls(db, project_id)
    if aps:
        stats = coverage_statistics(p.width_m, p.height_m, aps, walls)
        for k, v in stats["rssi"].items():
            w.writerow([f"RSSI {k}", f"{v['percent']}%", f"{v['count']} ячеек"])
    db.add(AuditLog(user_id=user.id, action="export_csv",
                    entity="project", entity_id=project_id))
    db.commit()
    buf.seek(0)
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition":
                 f'attachment; filename="project_{project_id}_{int(datetime.utcnow().timestamp())}.csv"'},
    )


@router.get("/{project_id}/xlsx")
def export_xlsx(project_id: int, db: Session = Depends(get_db),
                user: User = Depends(require_role("admin", "analyst"))):
    p = db.query(Project).get(project_id)
    if not p:
        raise HTTPException(status_code=404, detail="Проект не найден")
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="6C5CE7")
    title_font = Font(bold=True, size=14, color="2C2C3A")
    center = Alignment(horizontal="center", vertical="center")

    def style_header(ws, row=1):
        for cell in ws[row]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center

    def autosize(ws):
        for col in ws.columns:
            max_len = 0
            letter = get_column_letter(col[0].column)
            for cell in col:
                try:
                    v = len(str(cell.value)) if cell.value is not None else 0
                    max_len = max(max_len, v)
                except Exception:
                    pass
            ws.column_dimensions[letter].width = min(max_len + 4, 40)

    ws = wb.active
    ws.title = "Общая информация"
    ws["A1"] = f"Отчёт о проекте: {p.name}"
    ws["A1"].font = title_font
    ws["A3"] = "Размеры, м"
    ws["B3"] = f"{p.width_m} × {p.height_m}"
    ws["A4"] = "Площадь, м²"
    ws["B4"] = round(p.width_m * p.height_m, 1)
    ws["A5"] = "Дата создания"
    ws["B5"] = p.created_at.strftime("%d.%m.%Y %H:%M") if p.created_at else ""
    ws["A6"] = "Сформировал"
    ws["B6"] = f"{user.full_name} ({user.email})"
    autosize(ws)

    ws = wb.create_sheet("Точки доступа")
    ws.append(["ID", "Имя", "Модель", "X, м", "Y, м", "P, дБм",
               "G, dBi", "Диапазон", "Канал", "SSID"])
    style_header(ws)
    for ap in db.query(AccessPoint).filter(AccessPoint.project_id == project_id).all():
        ws.append([ap.id, ap.name, ap.model or "", ap.x, ap.y,
                   ap.tx_power_dbm, ap.antenna_gain, ap.band,
                   ap.channel, ap.ssid_type])
    autosize(ws)

    ws = wb.create_sheet("Устройства")
    ws.append(["ID", "Имя", "Тип", "X, м", "Y, м",
               "Требуемый RSSI", "Скорость, Мбит/с", "Диапазон"])
    style_header(ws)
    for d in db.query(ClientDevice).filter(ClientDevice.project_id == project_id).all():
        ws.append([d.id, d.name, d.type, d.x, d.y,
                   d.required_rssi, d.required_speed, d.band])
    autosize(ws)

    ws = wb.create_sheet("Стены")
    ws.append(["ID", "Тип", "Материал", "X1", "Y1", "X2", "Y2"])
    style_header(ws)
    for el in db.query(PlanElement).filter(
        PlanElement.project_id == project_id,
        PlanElement.type.in_(["wall", "door", "window"])
    ).all():
        ws.append([el.id, el.type, el.material or "",
                   el.x1, el.y1, el.x2, el.y2])
    autosize(ws)

    ws = wb.create_sheet("Измерения")
    ws.append(["ID", "Имя", "X, м", "Y, м", "RSSI, дБм",
               "SNR, дБ", "Интерференция", "Режим", "Дата"])
    style_header(ws)
    for m in db.query(Measurement).filter(Measurement.project_id == project_id).all():
        ws.append([m.id, m.name or "", m.x, m.y, m.rssi,
                   m.snr or "", m.interference or "", m.mode,
                   m.measured_at.strftime("%d.%m.%Y %H:%M") if m.measured_at else ""])
    autosize(ws)

    ws = wb.create_sheet("SSID")
    ws.append(["ID", "Имя", "Тип", "Аутентификация", "Шифрование",
               "VLAN", "Изоляция", "Внутр. доступ", "Интернет",
               "Captive Portal", "Лимит"])
    style_header(ws)
    for s in db.query(SSIDProfile).filter(SSIDProfile.project_id == project_id).all():
        ws.append([s.id, s.name, s.type, s.auth_method, s.encryption,
                   s.vlan_id or "",
                   "да" if s.client_isolation else "нет",
                   "да" if s.access_to_internal else "нет",
                   "да" if s.access_to_internet else "нет",
                   "да" if s.captive_portal else "нет",
                   s.bandwidth_limit_mbps or ""])
    autosize(ws)

    aps = _load_aps(db, project_id)
    walls = _load_walls(db, project_id)
    if aps:
        stats = coverage_statistics(p.width_m, p.height_m, aps, walls)
        ws = wb.create_sheet("Статистика")
        ws.append(["Метрика", "Ячеек", "% площади"])
        style_header(ws)
        for k, v in stats["rssi"].items():
            ws.append([f"RSSI: {k}", v["count"], v["percent"]])
        for k, v in stats["snr"].items():
            ws.append([f"SNR: {k}", v["count"], v["percent"]])
        autosize(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    db.add(AuditLog(user_id=user.id, action="export_xlsx",
                    entity="project", entity_id=project_id))
    db.commit()
    fname = f"project_{project_id}_{int(datetime.utcnow().timestamp())}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


@router.get("/{project_id}/data")
def get_report_data(project_id: int, db: Session = Depends(get_db),
                    user: User = Depends(require_role("admin", "analyst"))):
    p = db.query(Project).get(project_id)
    if not p:
        raise HTTPException(status_code=404, detail="Проект не найден")
    owner = db.query(User).get(p.owner_id) if p.owner_id else None
    aps = _load_aps(db, project_id)
    walls = _load_walls(db, project_id)
    stats = coverage_statistics(p.width_m, p.height_m, aps, walls) if aps else None
    channels = analyze_channels(aps) if aps else None
    db.add(AuditLog(user_id=user.id, action="export_pdf",
                    entity="project", entity_id=project_id))
    db.commit()
    return {
        "project": {"id": p.id, "name": p.name,
                    "width_m": p.width_m, "height_m": p.height_m,
                    "created_at": p.created_at.isoformat() if p.created_at else None},
        "owner": {"name": owner.full_name, "email": owner.email} if owner else None,
        "generated_by": {"name": user.full_name, "email": user.email,
                         "role": user.role},
        "generated_at": datetime.utcnow().isoformat(),
        "coverage_stats": stats, "channels": channels,
    }